"""
Sample Document Generator for All Supported Formats (PDF, DOCX, PNG, JPG, TXT).
Generates authentic and tampered academic credentials for live demonstration.
"""
import io
from pathlib import Path
from PIL import Image, ImageDraw
import docx


def generate_pdf_bytes(title: str, recipient: str, details: list[str]) -> bytes:
    """Generate a clean, standard PDF 1.4 document without external C-dependencies."""
    lines = [
        "TRUONG DAI HOC SECUDOC - ZERO-TRUST ACADEMY",
        "====================================================",
        title.upper(),
        "",
        f"Ho va Ten: {recipient}",
        f"Ngay sinh: 15/08/2004  |  Gioi tinh: Nam",
    ] + details + [
        "",
        "Chung thuc boi: Van phong Hoc vu SecuDoc",
        "Xac thuc so: HE THONG XAC THUC ZERO-TRUST SECUDOC"
    ]
    
    stream = "BT /F1 12 Tf 50 720 Td 18 TL "
    for l in lines:
        safe = l.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream += f"({safe}) ' "
    stream += "ET"
    sb = stream.encode("latin1", "replace")
    
    obj1 = b"1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n"
    obj2 = b"2 0 obj\n<< /Type /Pages /Kids [3 0 R] /Count 1 >>\nendobj\n"
    obj3 = b"3 0 obj\n<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>\nendobj\n"
    obj4 = f"4 0 obj\n<< /Length {len(sb)} >>\nstream\n".encode("latin1") + sb + b"\nendstream\nendobj\n"
    obj5 = b"5 0 obj\n<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>\nendobj\n"
    
    header = b"%PDF-1.4\n"
    body = [obj1, obj2, obj3, obj4, obj5]
    offsets = []
    curr = len(header)
    for o in body:
        offsets.append(curr)
        curr += len(o)
    xref = f"xref\n0 {len(body)+1}\n0000000000 65535 f \n".encode("latin1")
    for off in offsets:
        xref += f"{off:010d} 00000 n \n".encode("latin1")
    trailer = f"trailer\n<< /Size {len(body)+1} /Root 1 0 R >>\nstartxref\n{curr}\n%%EOF\n".encode("latin1")
    return header + b"".join(body) + xref + trailer


def generate_docx_bytes(title: str, recipient: str, mssv: str, courses: list[tuple[str, str, str]], gpa_str: str) -> bytes:
    """Generate a formal academic transcript DOCX using python-docx."""
    doc = docx.Document()
    
    # Title & Header
    doc.add_heading("TRƯỜNG ĐẠI HỌC SECUDOC - PHÒNG ĐÀO TẠO", level=1)
    p_sub = doc.add_paragraph("HỆ THỐNG VĂN BẰNG XÁC THỰC MẬT MÃ SỐ CHUẨN ZERO-TRUST")
    p_sub.italic = True
    
    doc.add_heading(title, level=2)
    doc.add_paragraph(f"Sinh viên: {recipient}  |  MSSV: {mssv}")
    doc.add_paragraph(f"Ngành đào tạo: An Toàn Thông Tin & Mạng Máy Tính")
    
    # Table of Course Grades
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Tên Môn Học"
    hdr[1].text = "Số Tín Chỉ"
    hdr[2].text = "Điểm Tổng Kết"
    
    for course_name, credits, grade in courses:
        row = table.add_row().cells
        row[0].text = course_name
        row[1].text = str(credits)
        row[2].text = str(grade)
        
    doc.add_paragraph()
    doc.add_paragraph(f"Điểm Trung Bình Tích Lũy (GPA): {gpa_str}")
    doc.add_paragraph("Xếp loại tốt nghiệp: XUẤT SẮC")
    doc.add_paragraph("Văn bản đã được ký số bằng khóa RSA-2048 của Hội đồng Đào tạo.")
    
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def generate_certificate_png_bytes(title: str, recipient: str, reason: str, is_revoked: bool = False) -> bytes:
    """Generate a certificate PNG image using PIL with borders and university seal."""
    img = Image.new("RGB", (800, 520), color=(253, 251, 247))
    d = ImageDraw.Draw(img)
    
    # Gold & Navy Borders
    d.rectangle([15, 15, 785, 505], outline=(184, 134, 11), width=4)
    d.rectangle([25, 25, 775, 495], outline=(10, 37, 64), width=2)
    
    # Headers
    d.text((230, 50), "TRUONG DAI HOC SECUDOC", fill=(10, 37, 64))
    d.text((250, 80), "VIEN KHOA HOC VA AN NINH MANG", fill=(100, 116, 139))
    d.text((220, 130), title.upper(), fill=(184, 134, 11))
    
    # Body
    d.text((120, 190), "Chung nhan va trao tang cho hoc vien:", fill=(51, 65, 85))
    d.text((260, 220), recipient.upper(), fill=(15, 23, 42))
    d.text((120, 280), f"Noi dung: {reason}", fill=(71, 85, 105))
    d.text((120, 320), "He thong xac thuc so Zero-Trust SecuDoc chung thuc tinh toan ven.", fill=(100, 116, 139))
    d.text((120, 350), "Trang thai mat ma: RSA-2048 Canonical Payload Sealed", fill=(71, 85, 105))
    
    if is_revoked:
        # Watermark/Stamp indicating revoked
        d.ellipse([580, 330, 750, 470], outline=(239, 68, 68), width=4)
        d.text((615, 385), "DA THU HOI", fill=(239, 68, 68))
        d.text((620, 410), "REVOKED", fill=(239, 68, 68))
    else:
        # Official Red Seal Circle
        d.ellipse([600, 350, 740, 470], outline=(220, 38, 38), width=3)
        d.text((630, 395), "DA KY SO", fill=(220, 38, 38))
        d.text((628, 418), "SECUDOC", fill=(220, 38, 38))
        
    bio = io.BytesIO()
    img.save(bio, format="PNG")
    return bio.getvalue()


def generate_student_card_jpg_bytes(recipient: str, mssv: str, dept: str, valid_until: str) -> bytes:
    """Generate a student ID card JPG using PIL."""
    img = Image.new("RGB", (640, 400), color=(240, 244, 248))
    d = ImageDraw.Draw(img)
    
    # Outer Card Frame
    d.rectangle([10, 10, 630, 390], outline=(30, 58, 138), width=3)
    # Header Banner
    d.rectangle([10, 10, 630, 75], fill=(30, 58, 138))
    d.text((30, 22), "TRUONG DAI HOC SECUDOC", fill=(255, 255, 255))
    d.text((30, 45), "THE SINH VIEN / STUDENT IDENTIFICATION CARD", fill=(191, 219, 254))
    
    # Photo Placeholder
    d.rectangle([40, 110, 180, 280], outline=(148, 163, 184), fill=(226, 232, 240), width=2)
    d.text((65, 185), "[ ANH THE ]", fill=(100, 116, 139))
    
    # Student Details
    d.text((220, 120), "Ho va Ten:", fill=(100, 116, 139))
    d.text((220, 140), recipient.upper(), fill=(15, 23, 42))
    
    d.text((220, 180), "Ma so sinh vien (MSSV):", fill=(100, 116, 139))
    d.text((220, 200), mssv, fill=(30, 58, 138))
    
    d.text((220, 240), "Khoa / Chuyen nganh:", fill=(100, 116, 139))
    d.text((220, 260), dept, fill=(15, 23, 42))
    
    d.text((220, 300), f"Co gia tri den: {valid_until}", fill=(71, 85, 105))
    d.text((40, 345), "CHIP XAC THUC MAT MA SHA-256 / RSA-2048 TICH HOP", fill=(100, 116, 139))
    
    bio = io.BytesIO()
    img.save(bio, format="JPEG", quality=95)
    return bio.getvalue()
